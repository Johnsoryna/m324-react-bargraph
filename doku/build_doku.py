# -*- coding: utf-8 -*-
"""
Erzeugt die Abgabe-Dokumentation als Word-Datei (python-docx).

Aufruf:  python build_doku.py
Ergebnis: M324_React_Package_Johnny_Leonhardt.docx im gleichen Ordner.

Screenshots: Liegt unter screenshots/<name>.png eine Datei, wird sie eingebettet,
sonst erscheint ein gelber Platzhalter mit dem Dateinamen. Danach Skript erneut ausfuehren.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "screenshots"
OUT = HERE / "M324_React_Package_Johnny_Leonhardt.docx"

BLUE = RGBColor(0x2F, 0x54, 0x96)
LIGHT_BLUE = RGBColor(0x2E, 0x74, 0xB5)

doc = Document()

# ---------- Styles ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)

for name, size, color in (("Heading 1", 14, BLUE), ("Heading 2", 12, LIGHT_BLUE)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 10)
    st.paragraph_format.space_after = Pt(4)

for section in doc.sections:
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.2)

# ---------- Helper ----------

def shade(cell_or_par, fill="F2F2F2"):
    """Hintergrundfarbe fuer Absatz oder Tabellenzelle."""
    pr = cell_or_par._p.get_or_add_pPr() if hasattr(cell_or_par, "_p") else cell_or_par._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def p(text="", bold=False, italic=False, size=None, align=None, after=6):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    return par


def rich(parts, after=6):
    """parts: Liste aus str oder (str, 'code'|'b'|'i')."""
    par = doc.add_paragraph()
    for part in parts:
        if isinstance(part, tuple):
            text, kind = part
            run = par.add_run(text)
            if kind == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif kind == "b":
                run.bold = True
            elif kind == "i":
                run.italic = True
        else:
            par.add_run(part)
    par.paragraph_format.space_after = Pt(after)
    return par


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def bullets(items):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        if isinstance(it, list):
            for part in it:
                if isinstance(part, tuple):
                    r = par.add_run(part[0])
                    if part[1] == "code":
                        r.font.name = "Consolas"
                        r.font.size = Pt(10)
                    elif part[1] == "b":
                        r.bold = True
                else:
                    par.add_run(part)
        else:
            par.add_run(it)
        par.paragraph_format.space_after = Pt(2)


def code(text):
    lines = text.strip("\n").splitlines()
    for i, line in enumerate(lines):
        par = doc.add_paragraph()
        run = par.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        pf = par.paragraph_format
        pf.left_indent = Cm(0.5)
        pf.space_after = Pt(8 if i == len(lines) - 1 else 0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        shade(par)


def table(header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        shade(cell, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            run.font.size = Pt(10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def screenshot(name, caption, width_cm=15.5):
    """Bild einbetten, falls screenshots/<name>.png existiert, sonst Platzhalter."""
    path = next((SHOTS / f"{name}.{ext}" for ext in ("png", "jpg", "jpeg") if (SHOTS / f"{name}.{ext}").exists()), None)
    if path is not None:
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(f"[Screenshot einfuegen: screenshots/{name}.png]")
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    cap.paragraph_format.space_after = Pt(12)


# =====================================================================
# Titel
# =====================================================================
p("Modul 324 DevOps", bold=True, size=24, after=2)
p("React Projekt mit Package", size=14, after=4)
p("Johnny Leonhardt, 5IA23b, 11.09.2026", after=4)
p("Repositories:", after=2)
bullets([
    "React-Package (Quellcode): github.com/Johnsoryna/m324-react-bargraph",
    "Publiziertes Package: npmjs.com/package/@johnsoryna/mybargraph",
    "Konsument (eigenes Package + Package des Lernpartners): github.com/Johnsoryna/m324-react-consumer",
])
p("Umgebung: Windows 11, Node 22.16.0, npm 11.4.2, Vite 8.3, React 19.2, Git 2.48, IntelliJ IDEA 2024.3")

# =====================================================================
h1("Aufgabe: Recherche: React-Komponente als Package bereitstellen")
p("Ein React-Package ist ein normales npm-Package. Der Unterschied zu einer React-App: Es hat keine "
  "index.html und kein main.jsx, sondern einen Einstiegspunkt (src/index.js), der die Komponenten exportiert. "
  "Ein Bundler baut daraus fertige JavaScript-Dateien in dist/. Der Konsument bekommt nicht den Quellcode, "
  "sondern dieses Artefakt aus node_modules, genau wie beim Maven-Artefakt in der letzten Aufgabe das JAR.")
p("Diese Punkte muss ein React-Package erfüllen:")
table(
    ["Punkt", "Bedeutung"],
    [
        ["Einstiegspunkt src/index.js", "Exportiert, was das Package nach aussen anbietet (Default- und/oder Named-Exports)."],
        ["Build-Output dist/", "JSX ist bereits nach JavaScript übersetzt. ES-Modul für Vite/webpack, CommonJS für require() und Jest."],
        ["main / module / exports", "Sagen npm und dem Bundler, welche Datei aus dist/ beim Import geladen wird."],
        ["files", "Nur dist/ (plus README, package.json) landet im Package, nicht Quellcode, Demo oder node_modules."],
        ["peerDependencies", "react und react-dom werden nicht mitgeliefert, das Zielprojekt hat sie schon. Sonst gibt es zwei React-Instanzen und Hooks funktionieren nicht."],
        ["CSS", "Muss ins Bundle injiziert werden (vite-plugin-lib-inject-css). Sonst muss der Konsument das CSS separat importieren."],
        ["Version", "Semantic Versioning. Jede Publikation braucht eine neue, höhere Versionsnummer."],
    ],
    widths=[4.5, 11.5],
)
p("Zum Bauen gibt es mehrere Werkzeuge. Das Medium-Tutorial aus dem Auftrag ist hinter einer Paywall, "
  "sichtbar ist nur die Einleitung. Der nPlan-Artikel (Rollup + GitHub Packages) ist vollständig lesbar und "
  "zeigt den Rollup-Weg.")
table(
    ["Werkzeug", "Vorgehen", "Vorteile", "Nachteile"],
    [
        ["Vite Library-Modus", "build.lib in vite.config.js, Formate es + cjs, react als external",
         "Gleiche Toolchain wie die App, wenig Konfiguration, Dev-Server für eine Demo inklusive",
         "Ein Package pro Projekt, mehrere Packages aus einem Repo sind mühsam"],
        ["Rollup direkt (nPlan-Artikel)", "rollup.config.js mit Babel, commonjs, node-resolve, postcss, peer-deps-external",
         "Volle Kontrolle, mehrere Packages aus einem Repo möglich",
         "Viele Plugins, viel Konfiguration, Babel separat einrichten"],
        ["tsup / esbuild", "Ein Befehl: tsup src/index.js --format esm,cjs",
         "Sehr schnell, fast keine Konfiguration",
         "CSS-Handling schwach, eher für TypeScript-Libraries ohne Styles"],
        ["Ohne Build (JSX publizieren)", "Quellcode direkt in files aufnehmen",
         "Nichts zu bauen",
         "Konsument muss JSX aus node_modules selbst kompilieren, das tun Vite/webpack standardmässig nicht. Unüblich."],
    ],
    widths=[3.2, 4.6, 4.1, 4.1],
)
p("Entscheidung: Vite Library-Modus. Das Projekt ist mit Vite erstellt, die Anleitung im Modul baut darauf auf, "
  "und der Dev-Server lässt sich für eine Demo der Komponente weiterverwenden.")

# =====================================================================
h1("Aufgabe: Recherche: Alternative: publish on GitHub")
p("GitHub Packages enthält neben der Maven-Registry auch eine npm-Registry unter npm.pkg.github.com. "
  "Das Vorgehen ist das gleiche wie bei GitHub Maven Packages in der letzten Aufgabe: Token mit "
  "write:packages, Registry-URL angeben, publizieren. Die Unterschiede zu npmjs.com liegen im Detail:")
bullets([
    ["Der Package-Name muss mit dem GitHub-Benutzer als Scope beginnen: ", ("@johnsoryna/mybargraph", "code"), "."],
    ["Eine ", (".npmrc", "code"), " im Projekt verbindet den Scope mit der Registry und liefert das Token. "
     "Alternativ ", ("publishConfig.registry", "code"), " in der package.json."],
    ["Das Feld ", ("repository.url", "code"), " in der package.json muss auf das GitHub-Repository zeigen, "
     "sonst lehnt die Registry das Package ab."],
    ["Der Token ist ein Personal Access Token (classic) mit ", ("write:packages", "code"), " (enthält ",
     ("read:packages", "code"), ") und ", ("repo", "code"), "."],
    ["Auch der Konsument braucht eine ", (".npmrc", "code"), " mit Scope-Mapping und einem Token mit ",
     ("read:packages", "code"), ", sogar bei öffentlichen Packages. Das ist der grösste Unterschied zu npmjs.com."],
])
p("Konfiguration für GitHub Packages (nicht umgesetzt, nur recherchiert):")
code("""
# .npmrc (Package-Projekt und Konsument)
@johnsoryna:registry=https://npm.pkg.github.com
//npm.pkg.github.com/:_authToken=${GITHUB_TOKEN}

# package.json (Package-Projekt)
"publishConfig": { "registry": "https://npm.pkg.github.com" },
"repository": { "type": "git", "url": "git+https://github.com/Johnsoryna/m324-react-bargraph.git" }

npm publish
""")
table(
    ["Kriterium", "npmjs.com", "GitHub Packages"],
    [
        ["Sichtbarkeit", "Öffentlich für alle, private Packages kostenpflichtig", "Public oder private, an Repository gebunden"],
        ["Installation beim Konsument", "npm install, keine Konfiguration", ".npmrc mit Scope und Token nötig, auch bei public"],
        ["Authentifizierung beim Publizieren", "npm login (Browser) oder Access-Token, bei 2FA zusätzlich OTP", "Personal Access Token mit write:packages"],
        ["Kosten", "Gratis für public", "Gratis für public, Speicher- und Traffic-Limit bei private"],
        ["Versionen löschen", "Version unveränderlich, unpublish nur innerhalb 72 h", "Versionen können gelöscht werden"],
        ["CI/CD", "Trusted Publishing (OIDC) oder Token als Secret", "GITHUB_TOKEN in Actions direkt nutzbar"],
        ["Eignung für Lernpartner", "Einfach, jeder kann sofort installieren", "Jeder Partner braucht einen GitHub-Token"],
    ],
    widths=[4.0, 6.0, 6.0],
)
p("Entscheidung: npmjs.com. Die Lernpartner können das Package ohne Token installieren. GitHub Packages "
  "passt, wenn ein Package privat bleiben soll und alle Konsumenten sowieso in derselben GitHub-Organisation "
  "sind, wie im nPlan-Artikel beschrieben.")

# =====================================================================
h1("Aufgabe: React-Package erstellen")
p("Ziel: Die Komponente HorizontalBarGraph als npm-Package @johnsoryna/mybargraph, gebaut mit Vite im Library-Modus.")

h2("Schritt 1: Projekt anlegen und Template ausmisten")
p("Vite kann kein Library-Projekt direkt erzeugen. Darum zuerst ein normales React-Projekt "
  "(IntelliJ: File > New > Project > Vite, Framework React, JavaScript; im Terminal gleichwertig):")
code("npm create vite@latest mybargraph -- --template react\ncd mybargraph\nnpm install")
p("Danach alles entfernt, was nur eine App braucht: index.html, public/, src/App.jsx, src/App.css, "
  "src/index.css, src/main.jsx, src/assets. Übrig bleibt diese Struktur:")
code("""
mybargraph/
    package.json               Name, Version, files, exports, peerDependencies
    vite.config.js             Library-Modus (Build nach dist/)
    vite.demo.config.js        nur fuer die lokale Demo
    src/index.js               Einstiegspunkt, exportiert die Komponente
    src/components/HorizontalBarGraph.jsx
    src/components/HorizontalBarGraph.css
    demo/index.html, demo/src/main.jsx   Demo-App (nicht im Package)
    dist/                      Build-Ergebnis (in .gitignore, aber im Package)
""")
screenshot("intellij_projekt", "Projektstruktur des Packages in IntelliJ")

h2("Schritt 2: vite.config.js auf Library-Modus umstellen")
code("""
import { defineConfig } from 'vite'
import react from '@vitejs/plugin-react'
import { fileURLToPath } from 'node:url'
import { libInjectCss } from 'vite-plugin-lib-inject-css'

export default defineConfig({
  plugins: [react(), libInjectCss()],
  build: {
    lib: {
      entry: fileURLToPath(new URL('./src/index.js', import.meta.url)),
      name: 'MyBargraph',
      formats: ['es', 'cjs'],
      fileName: (format) => `mybargraph.${format}.js`,
    },
    rollupOptions: {
      external: ['react', 'react-dom', 'react/jsx-runtime', 'prop-types'],
      output: { exports: 'named', globals: { react: 'React', 'react-dom': 'ReactDOM' } },
    },
  },
})
""")
p("Zwei Abweichungen von der Anleitung im Modul: Die package.json hat \"type\": \"module\", darum gibt es "
  "__dirname nicht, ich verwende fileURLToPath(new URL(...)). Und exports: 'named' verhindert die Warnung "
  "MIXED_EXPORTS, weil src/index.js sowohl einen Default- als auch einen Named-Export hat. "
  "libInjectCss schreibt einen import './index.css' in das gebaute Modul, so lädt der Konsument das CSS automatisch mit.")

h2("Schritt 3: package.json anpassen")
code("""
{
  "name": "@johnsoryna/mybargraph",
  "version": "1.0.0",
  "private": false,
  "type": "module",
  "license": "MIT",
  "repository": { "type": "git", "url": "git+https://github.com/Johnsoryna/m324-react-bargraph.git" },
  "files": ["dist"],
  "main": "./dist/mybargraph.cjs.js",
  "module": "./dist/mybargraph.es.js",
  "exports": {
    ".": { "import": "./dist/mybargraph.es.js", "require": "./dist/mybargraph.cjs.js" }
  },
  "sideEffects": ["**/*.css"],
  "scripts": { "build": "vite build", "demo": "vite --config vite.demo.config.js" },
  "peerDependencies": { "react": ">=18", "react-dom": ">=18" },
  "dependencies": { "prop-types": "^15.8.1" },
  "devDependencies": { "react": "^19.2.8", "react-dom": "^19.2.8", "vite": "^8.3.0",
                       "@vitejs/plugin-react": "^6.1.1", "vite-plugin-lib-inject-css": "^2.2.2" }
}
""")
bullets([
    ["Der Name trägt den npm-Benutzernamen als Scope (", ("@johnsoryna/", "code"), "), damit er eindeutig ist."],
    ["react und react-dom stehen in ", ("peerDependencies", "code"), " (der Konsument liefert sie) und zusätzlich in ",
     ("devDependencies", "code"), ", damit Build und Demo lokal funktionieren."],
    [("files: [\"dist\"]", "code"), " sorgt dafür, dass Quellcode und Demo nicht im Package landen."],
    [("sideEffects: [\"**/*.css\"]", "code"), " verhindert, dass ein Bundler den CSS-Import beim Tree-Shaking entfernt."],
])

h2("Schritt 4 und 5: Einstiegspunkt und Komponente")
code("""
// src/index.js
export { default } from './components/HorizontalBarGraph.jsx'
export { default as HorizontalBarGraph } from './components/HorizontalBarGraph.jsx'
""")
code("""
// src/components/HorizontalBarGraph.jsx (gekuerzt)
import PropTypes from 'prop-types'
import './HorizontalBarGraph.css'

export default function HorizontalBarGraph({ value, maxvalue = 100, barwidthpx = 300,
                                             label = '', color = '#2563eb', showValue = true }) {
  const safeMax = maxvalue > 0 ? maxvalue : 1
  const clamped = Math.min(Math.max(value, 0), safeMax)
  const percent = Math.round((clamped / safeMax) * 100)
  return (
    <div className="hbg" style={{ width: `${barwidthpx}px` }}>
      <div className="hbg__header"><span>{label}</span>
        {showValue && <span>{clamped} / {safeMax} ({percent}%)</span>}</div>
      <div className="hbg__track" role="progressbar" aria-valuenow={clamped} aria-valuemax={safeMax}>
        <div className="hbg__bar" style={{ width: `${percent}%`, backgroundColor: color }} />
      </div>
    </div>
  )
}
HorizontalBarGraph.propTypes = { value: PropTypes.number.isRequired, maxvalue: PropTypes.number, ... }
""")
p("Der Balken wächst per CSS-Transition auf die neue Breite, wenn sich value ändert. Props der Komponente:")
table(
    ["Prop", "Typ", "Default", "Beschreibung"],
    [
        ["value", "number", "(Pflicht)", "Aktueller Wert, wird auf 0..maxvalue begrenzt"],
        ["maxvalue", "number", "100", "Maximalwert, entspricht 100 %"],
        ["barwidthpx", "number", "300", "Gesamtbreite des Balkens in Pixel"],
        ["label", "string", "''", "Beschriftung über dem Balken"],
        ["color", "string", "'#2563eb'", "Farbe des gefüllten Balkens"],
        ["showValue", "boolean", "true", "Wert und Prozentzahl rechts anzeigen"],
    ],
    widths=[3.0, 2.2, 2.6, 8.2],
)

h2("Schritt 6 und 7: prop-types installieren und bauen")
code("""
npm install prop-types --save
npm install -D vite-plugin-lib-inject-css
npm run build

dist/index.css           0.51 kB
dist/mybargraph.es.js    1.32 kB
dist/mybargraph.cjs.js   1.57 kB
built in 126ms
""")
p("Die Bundles sind klein, weil react, react-dom und prop-types nicht enthalten sind. "
  "npm pack --dry-run zeigt, was tatsächlich publiziert würde:")
code("""
npm notice package: @johnsoryna/mybargraph@1.0.0
npm notice 1.7kB README.md
npm notice 514B  dist/index.css
npm notice 1.6kB dist/mybargraph.cjs.js
npm notice 1.3kB dist/mybargraph.es.js
npm notice 1.3kB package.json
npm notice package size: 3.0 kB, total files: 5
""")
screenshot("term_build", "npm run build und npm pack --dry-run im Terminal")

h2("Demo im Package selbst")
p("Für die Demo gibt es den Ordner demo/ mit index.html und src/main.jsx, die die Komponente direkt aus "
  "src/ importiert. Das Vorgehen aus der Anleitung (\"demo\": \"vite demo\") hat bei mir nicht funktioniert: "
  "Vite sucht die vite.config.js im angegebenen Root, also in demo/, findet sie nicht und hat dann kein "
  "React-Plugin. Darum eine eigene kleine Konfiguration:")
code("""
// vite.demo.config.js
export default defineConfig({ root: 'demo', plugins: [react()] })

// package.json
"demo": "vite --config vite.demo.config.js"
""")
screenshot("demo_browser", "npm run demo: Demo-Seite mit Slider und drei Balken im Browser")

# =====================================================================
h1("Aufgabe: Package publizieren (npmjs.com)")
p("Voraussetzung ist ein Account auf npmjs.com, der Benutzername (johnsoryna) ist der Scope im Package-Namen. "
  "Für die Authentifizierung gibt es zwei Wege:")
table(
    ["Weg", "Vorgehen", "Bewertung"],
    [
        ["npm login", "Öffnet den Browser, man meldet sich an, das Token landet in ~/.npmrc des Benutzers",
         "Kein Token im Projekt, für die manuelle Abgabe ideal"],
        ["Access-Token (Anleitung im Modul)", "Token auf npmjs.com erzeugen, in .npmrc im Projekt eintragen, .npmrc in .gitignore",
         "Nötig für CI/CD; Risiko, dass das Token versehentlich committet wird"],
    ],
    widths=[3.5, 7.0, 5.5],
)
p("Ich habe npm login gewählt. Die Datei .npmrc steht trotzdem in .gitignore, falls später ein Token für eine "
  "Pipeline dazukommt. Publiziert wird mit --access public, weil Packages mit Scope standardmässig privat wären "
  "und private Packages auf npmjs.com kostenpflichtig sind.")
p("Der erste Versuch scheiterte mit 403: npmjs.com verlangt seit 2025 für jedes Publizieren eine "
  "Zwei-Faktor-Authentifizierung, Tokens mit \"Bypass 2FA\" (wie in der Modul-Anleitung beschrieben) werden "
  "eingeschränkt. Nach dem Aktivieren von 2FA mit einem Passkey öffnet npm publish einen Browser-Link, dort "
  "bestätigt man mit dem Passkey, und das Terminal fährt fort. Wichtig: Das funktioniert nur in einem "
  "interaktiven Terminal, nicht in einem Skript.")
code("""
PS C:\Code\324\mybargraph> npm login
Logged in on https://registry.npmjs.org/.

PS C:\Code\324\mybargraph> npm publish --access public
npm notice package: @johnsoryna/mybargraph@1.0.0
npm notice total files: 5
npm notice Publishing to https://registry.npmjs.org/ with tag latest and public access
Authenticate your account at:
https://www.npmjs.com/auth/cli/984ba77b-....
Press ENTER to open in the browser...

+ @johnsoryna/mybargraph@1.0.0
""")
screenshot("term_publish", "npm publish --access public im Terminal")
screenshot("npmjs_package", "Das Package @johnsoryna/mybargraph auf npmjs.com")
p("Eine publizierte Version ist unveränderlich. Für jede weitere Publikation muss die Version erhöht werden, "
  "z.B. mit npm version patch (1.0.0 -> 1.0.1), erst dann geht npm publish wieder durch.")

# =====================================================================
h1("Aufgabe: Package im eigenen Projekt verwenden")
p("Zweites Projekt bargraph-consumer, ebenfalls mit Vite und React. Das Package wird wie jede andere "
  "Abhängigkeit installiert und landet unter node_modules/@johnsoryna/mybargraph, nur mit dist/, README "
  "und package.json, ohne Quellcode.")
code("""
npm create vite@latest bargraph-consumer -- --template react
cd bargraph-consumer
npm install
npm install @johnsoryna/mybargraph
""")
code("""
// src/App.jsx (gekuerzt)
import { useState } from 'react'
import HorizontalBarGraph from '@johnsoryna/mybargraph'

function App() {
  const [value, setValue] = useState(42)
  return (
    <>
      <input type="range" min="0" max="100" value={value}
             onChange={(e) => setValue(Number(e.target.value))} />
      <HorizontalBarGraph value={value} maxvalue={100} barwidthpx={420} label="Fortschritt" />
      <HorizontalBarGraph value={value} maxvalue={60} barwidthpx={300} label="Ziel 60" color="#16a34a" />
    </>
  )
}
""")
p("In der package.json des Konsumenten steht danach \"@johnsoryna/mybargraph\": \"^1.0.0\". Das CSS muss "
  "nicht separat importiert werden, weil es im Bundle injiziert ist.")
screenshot("explorer_node_modules", "node_modules/@johnsoryna/mybargraph im Konsumenten-Projekt: nur dist/, README und package.json")
screenshot("consumer_browser", "npm run dev im Konsumenten: HorizontalBarGraph aus dem npm-Package")

# =====================================================================
h1("Aufgabe: Package vom Lernpartner einbinden")
par = p("Lernpartner: [NAME]   Package: [@PARTNER/PAKETNAME]   Registry: [npmjs.com / GitHub Packages]", bold=True)
for run in par.runs:
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
p("Das Package des Lernpartners wird im gleichen Konsumenten-Projekt installiert und in src/PartnerDemo.jsx "
  "verwendet. Die Datei enthält bereits die vorbereitete Stelle:")
code("""
npm install @PARTNER/PAKETNAME

// src/PartnerDemo.jsx
import PartnerKomponente from '@PARTNER/PAKETNAME'
...
<PartnerKomponente prop1={...} prop2={...} />
""")
p("Falls der Lernpartner auf GitHub Packages publiziert hat, braucht es zusätzlich eine .npmrc im Konsumenten "
  "mit dem Scope-Mapping (@partner:registry=https://npm.pkg.github.com) und einem eigenen GitHub-Token mit "
  "read:packages, siehe Recherche oben.")
screenshot("partner_browser", "Konsument mit eigenem Package und dem Package des Lernpartners")
screenshot("partner_package_json", "package.json des Konsumenten mit beiden Abhängigkeiten")

# =====================================================================
h1("Fazit und Probleme")
bullets([
    "Ein React-Package unterscheidet sich von einer App nur durch Einstiegspunkt, Library-Build und die "
    "package.json-Felder files, exports und peerDependencies. Der Rest ist normales React.",
    "__dirname aus der Anleitung existiert in einer ESM-Konfiguration nicht, fileURLToPath(new URL(...)) ersetzt es.",
    "\"vite demo\" findet die Konfiguration nicht, darum eine separate vite.demo.config.js.",
    "prop-types als Dependency, nicht als Peer, und trotzdem external im Build: npm installiert es beim Konsumenten mit, "
    "aber es wird nicht ins Bundle kopiert.",
    "Der Unterschied zu Maven: Bei npm ist der Konsument-Registry-Zugriff ohne Konfiguration möglich, bei GitHub "
    "Packages braucht auch npm (wie Maven) ein Token zum Lesen.",
    "npmjs.com erzwingt 2FA beim Publizieren. Der Access-Token-Weg mit \"Bypass 2FA\" aus der Anleitung wird "
    "eingeschränkt, für Pipelines ist Trusted Publishing (OIDC) der vorgesehene Ersatz.",
])

# =====================================================================
doc.add_page_break()
h1("Anhang: Checkliste Screenshots (vor der PDF-Abgabe entfernen)")
p("Dateien nach doku/screenshots/ legen und python build_doku.py erneut ausführen, dann werden sie automatisch eingebettet.")
table(
    ["Datei", "Inhalt", "Wo"],
    [
        ["intellij_projekt.png", "IntelliJ mit Projektbaum von mybargraph (src, demo, dist, package.json)", "Package erstellen, Schritt 1"],
        ["term_build.png", "Terminal: npm run build und npm pack --dry-run", "Package erstellen, Schritt 6/7"],
        ["demo_browser.png", "Browser: npm run demo mit Slider und Balken", "Demo"],
        ["term_publish.png", "Terminal: npm whoami und npm publish --access public", "Publizieren"],
        ["npmjs_package.png", "Browser: npmjs.com/package/@johnsoryna/mybargraph", "Publizieren"],
        ["explorer_node_modules.png", "Explorer/IntelliJ: node_modules/@johnsoryna/mybargraph/dist im Konsumenten", "Verwenden"],
        ["consumer_browser.png", "Browser: bargraph-consumer mit npm run dev", "Verwenden"],
        ["partner_browser.png", "Browser: Konsument mit Partner-Komponente", "Lernpartner"],
        ["partner_package_json.png", "package.json des Konsumenten mit beiden Packages", "Lernpartner"],
    ],
    widths=[4.5, 7.5, 4.0],
)

doc.save(OUT)
print(f"geschrieben: {OUT}")
